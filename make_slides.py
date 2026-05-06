from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import pptx.oxml.ns as nsmap
from lxml import etree

# ── Brand colours ──────────────────────────────────────────────
INDIGO       = RGBColor(0x4F, 0x46, 0xE5)
INDIGO_DARK  = RGBColor(0x37, 0x30, 0xA3)
INDIGO_LIGHT = RGBColor(0xEE, 0xF2, 0xFF)
SLATE        = RGBColor(0x47, 0x55, 0x69)
SLATE_LIGHT  = RGBColor(0xF1, 0xF5, 0xF9)
EMERALD      = RGBColor(0x05, 0x96, 0x69)
EMERALD_LIGHT= RGBColor(0xD1, 0xFA, 0xE5)
AMBER        = RGBColor(0xD9, 0x77, 0x06)
AMBER_LIGHT  = RGBColor(0xFE, 0xF3, 0xC7)
ROSE         = RGBColor(0xE1, 0x1D, 0x48)
ROSE_LIGHT   = RGBColor(0xFF, 0xE4, 0xE6)
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
NEAR_WHITE   = RGBColor(0xF8, 0xFA, 0xFF)
DARK_TEXT    = RGBColor(0x1E, 0x1B, 0x4B)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

def blank_slide(prs):
    blank = prs.slide_layouts[6]
    return prs.slides.add_slide(blank)

def bg_rect(slide, color, left=0, top=0, width=None, height=None):
    if width is None:  width  = prs.slide_width
    if height is None: height = prs.slide_height
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_textbox(slide, text, left, top, width, height,
                font_size=14, bold=False, color=DARK_TEXT,
                align=PP_ALIGN.LEFT, italic=False, wrap=True):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return txBox

def add_para(tf, text, font_size=11, bold=False, color=DARK_TEXT,
             align=PP_ALIGN.LEFT, space_before=0, italic=False):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return p

def pill(slide, text, left, top, width, height, bg, fg=WHITE, font_size=9):
    shape = slide.shapes.add_shape(
        pptx.util.MSO_SHAPE_TYPE if False else 5,  # rounded rect = 5
        left, top, width, height
    )
    # Use rounded rectangle (freeform won't work easily, use built-in rounded rect)
    from pptx.oxml.ns import qn
    sp = shape._element
    prstGeom = sp.find(qn('p:spPr')).find(qn('a:prstGeom'))
    if prstGeom is not None:
        prstGeom.set('prst', 'roundRect')
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = True
    run.font.color.rgb = fg
    return shape

def colored_bar(slide, left, top, height, color, width=Inches(0.06)):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def section_header(slide, title, left, top, width, color=INDIGO_DARK, font_size=13):
    txBox = slide.shapes.add_textbox(left, top, width, Inches(0.35))
    tf = txBox.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(font_size)
    run.font.bold = True
    run.font.color.rgb = color
    return txBox

# ══════════════════════════════════════════════════════════════════
#  SLIDE 1 — INTRO
# ══════════════════════════════════════════════════════════════════
s1 = blank_slide(prs)

# Dark indigo background
bg_rect(s1, INDIGO_DARK)

# Top accent strip
bg_rect(s1, INDIGO, left=0, top=0, width=prs.slide_width, height=Inches(0.07))

# Logo area — text logo
add_textbox(s1, "EduReach", Inches(0.6), Inches(0.5), Inches(5), Inches(0.9),
            font_size=42, bold=True, color=WHITE)

# Tagline
add_textbox(s1, "Exam prep & academic competition platform for African students",
            Inches(0.6), Inches(1.5), Inches(8), Inches(0.55),
            font_size=16, color=RGBColor(0xC7, 0xD2, 0xFE))

# Thin rule
rule = s1.shapes.add_shape(1, Inches(0.6), Inches(2.25), Inches(12.1), Pt(1.2))
rule.fill.solid()
rule.fill.fore_color.rgb = RGBColor(0x6B, 0x73, 0xFF)
rule.line.fill.background()

# ── 4 stat pills ──────────────────────────────────────────────
stats = [
    ("117", "USERS"),
    ("33%", "ACTIVE"),
    ("MVP", "COMPLETE"),
    ("KES 0", "REVENUE"),
]
pill_w, pill_h = Inches(2.5), Inches(1.1)
gap = Inches(0.25)
start_x = Inches(0.6)
for i, (val, label) in enumerate(stats):
    x = start_x + i * (pill_w + gap)
    y = Inches(2.55)
    box = s1.shapes.add_shape(5, x, y, pill_w, pill_h)
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x3B, 0x35, 0xB8)
    box.line.fill.background()
    tf = box.text_frame
    tf.word_wrap = False
    # value
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = val
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = WHITE
    # label
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(2)
    r2 = p2.add_run()
    r2.text = label
    r2.font.size = Pt(8)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0xA5, 0xB4, 0xFC)

# ── What is EduReach? ─────────────────────────────────────────
y_section = Inches(3.95)
add_textbox(s1, "What is EduReach?", Inches(0.6), y_section, Inches(12), Inches(0.4),
            font_size=13, bold=True, color=RGBColor(0xA5, 0xB4, 0xFC))

bullets = [
    ("AI-powered exam preparation", "Practice papers, MCQs, essay grading, and olympiad challenges — all in one platform."),
    ("Built for African students", "Aligned with KCSE, A-Levels, IB, university entrance exams, and maths olympiads."),
    ("Study groups & live challenges", "Students compete, collaborate, and track progress together in subject groups."),
    ("Guest trial — no account needed", "14-day free trial with no sign-up friction. Users experience value before committing."),
]

y = Inches(4.4)
for title, desc in bullets:
    # Dot
    dot = s1.shapes.add_shape(9, Inches(0.6), y + Inches(0.06), Inches(0.12), Inches(0.12))
    dot.fill.solid()
    dot.fill.fore_color.rgb = INDIGO
    dot.line.fill.background()
    # Text
    txb = s1.shapes.add_textbox(Inches(0.85), y, Inches(11.8), Inches(0.38))
    tf = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = title + "  "
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = WHITE
    r2 = p.add_run()
    r2.text = desc
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor(0xC7, 0xD2, 0xFE)
    y += Inches(0.47)

# Footer
add_textbox(s1, "Dalton Omondi & Mary Syokau  |  Nairobi, Kenya  |  May 2026",
            Inches(0.6), Inches(7.05), Inches(12), Inches(0.35),
            font_size=9, color=RGBColor(0x6B, 0x73, 0xFF), align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════
#  SLIDE 2 — WEEK 1 DELIVERABLES & CHALLENGES
# ══════════════════════════════════════════════════════════════════
s2 = blank_slide(prs)
bg_rect(s2, NEAR_WHITE)

# Header bar
bg_rect(s2, INDIGO_DARK, left=0, top=0, width=prs.slide_width, height=Inches(0.9))
add_textbox(s2, "Week 1 — What We Shipped & Challenges We Faced",
            Inches(0.4), Inches(0.15), Inches(11), Inches(0.6),
            font_size=20, bold=True, color=WHITE)
add_textbox(s2, "Features shipped  ·  Bugs fixed  ·  Technical blockers  ·  Growth & financial pressure",
            Inches(0.4), Inches(0.6), Inches(11), Inches(0.32),
            font_size=10, color=RGBColor(0xC7, 0xD2, 0xFE))

# ── LEFT: What We Shipped ──────────────────────────────────────
col_left = Inches(0.35)
col_right = Inches(6.85)
col_w = Inches(6.1)
y0 = Inches(1.05)

section_header(s2, "What We Shipped", col_left, y0, col_w, color=INDIGO_DARK)

shipped = [
    (EMERALD, "Guest trial — 14 days, no account needed"),
    (EMERALD, "Deferred AI marking (submit now, grade on demand)"),
    (EMERALD, "Image answer upload (photo of written work)"),
    (EMERALD, "Study group searchable challenge picker"),
    (EMERALD, "Personalised dashboard with relevance ranking"),
    (EMERALD, "Pop-out AI tutor accessible on all pages"),
    (EMERALD, "Olympiad papers seeded — multi-subject"),
    (EMERALD, "Mobile tab overflow fixed (Events & Invites)"),
    (EMERALD, "Google Sign-In fixed (switched to renderButton)"),
    (EMERALD, "Assessment list cap removed (was stuck at 20)"),
    (EMERALD, "Clean error pages in production"),
    (EMERALD, "Microeconomics — 6 papers, 60 real MCQ questions"),
]

y = y0 + Inches(0.42)
for color, text in shipped:
    # Green check pill
    ck = s2.shapes.add_shape(5, col_left, y + Inches(0.02), Inches(0.55), Inches(0.22))
    ck.fill.solid()
    ck.fill.fore_color.rgb = EMERALD
    ck.line.fill.background()
    tf_ck = ck.text_frame
    p_ck = tf_ck.paragraphs[0]
    p_ck.alignment = PP_ALIGN.CENTER
    r_ck = p_ck.add_run()
    r_ck.text = "Done"
    r_ck.font.size = Pt(6.5)
    r_ck.font.bold = True
    r_ck.font.color.rgb = WHITE

    txb = s2.shapes.add_textbox(col_left + Inches(0.62), y, col_w - Inches(0.65), Inches(0.27))
    tf = txb.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(9.5)
    run.font.color.rgb = DARK_TEXT
    y += Inches(0.29)

# ── RIGHT: Challenges ──────────────────────────────────────────
section_header(s2, "Challenges We Faced", col_right, y0, col_w, color=INDIGO_DARK)

challenges = [
    (ROSE,   "CRITICAL",  "Cash position — KES 144 left after committed costs. Zero buffer."),
    (ROSE,   "CRITICAL",  "67% of users sign up and never return — no first-session value."),
    (ROSE,   "CRITICAL",  "Zero paying users despite billing being live."),
    (AMBER,  "TECH",      "Render → Contabo migration mid-sprint (RAM limit crashes)."),
    (AMBER,  "TECH",      "Guest 401 errors crashing the app for all unauthenticated visitors."),
    (AMBER,  "TECH",      "Google Sign-In silently failing in all browsers."),
    (AMBER,  "GROWTH",    "No upgrade prompt shown after sign-up."),
    (SLATE,  "PENDING",   "African Olympiad MOU — not yet signed."),
    (SLATE,  "PENDING",   "Kenyan IMO team onboarding not yet started."),
    (SLATE,  "CONTENT",   "Placeholder exam content in database (now fixed)."),
]

y = y0 + Inches(0.42)
for color, tag, text in challenges:
    # Tag pill
    tag_w = Inches(0.72)
    tp = s2.shapes.add_shape(5, col_right, y + Inches(0.02), tag_w, Inches(0.22))
    tp.fill.solid()
    tp.fill.fore_color.rgb = color
    tp.line.fill.background()
    tf_tp = tp.text_frame
    p_tp = tf_tp.paragraphs[0]
    p_tp.alignment = PP_ALIGN.CENTER
    r_tp = p_tp.add_run()
    r_tp.text = tag
    r_tp.font.size = Pt(5.5)
    r_tp.font.bold = True
    r_tp.font.color.rgb = WHITE

    txb = s2.shapes.add_textbox(col_right + Inches(0.78), y, col_w - Inches(0.82), Inches(0.27))
    tf = txb.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(9.5)
    run.font.color.rgb = DARK_TEXT
    y += Inches(0.29)

# Divider line between columns
div = s2.shapes.add_shape(1, Inches(6.7), Inches(1.0), Pt(1), Inches(6.3))
div.fill.solid()
div.fill.fore_color.rgb = RGBColor(0xC7, 0xD2, 0xFE)
div.line.fill.background()

# Footer
bg_rect(s2, INDIGO_DARK, left=0, top=Inches(7.18), width=prs.slide_width, height=Inches(0.32))
add_textbox(s2, "EduReach  |  Week 1 Review  |  May 2026",
            Inches(0.4), Inches(7.2), Inches(7), Inches(0.25),
            font_size=8, color=RGBColor(0xA5, 0xB4, 0xFC))
add_textbox(s2, "Slide 2 / 3",
            Inches(11.5), Inches(7.2), Inches(1.5), Inches(0.25),
            font_size=8, color=RGBColor(0xA5, 0xB4, 0xFC), align=PP_ALIGN.RIGHT)


# ══════════════════════════════════════════════════════════════════
#  SLIDE 3 — WEEK 2 DELIVERABLES
# ══════════════════════════════════════════════════════════════════
s3 = blank_slide(prs)
bg_rect(s3, NEAR_WHITE)

# Header bar
bg_rect(s3, INDIGO_DARK, left=0, top=0, width=prs.slide_width, height=Inches(0.9))
add_textbox(s3, "Week 2 — Deliverables",
            Inches(0.4), Inches(0.15), Inches(11), Inches(0.6),
            font_size=20, bold=True, color=WHITE)
add_textbox(s3, "What must happen in the next 7 days — ranked by urgency",
            Inches(0.4), Inches(0.6), Inches(11), Inches(0.32),
            font_size=10, color=RGBColor(0xC7, 0xD2, 0xFE))

# ── LEFT: Must Ship ────────────────────────────────────────────
y0 = Inches(1.05)
section_header(s3, "Must Ship", col_left, y0, Inches(6.1), color=ROSE, font_size=13)

must_ship = [
    (ROSE,  "1", "Secure first paying customer",
     "One school, teacher, or parent — any amount. Billing + M-Pesa are live. This is a sales call, not a tech task."),
    (ROSE,  "2", "Activate the upgrade funnel",
     "Plan selection screen (Free / Starter / Pro) shown immediately after sign-up. Users must be prompted, not left to discover it."),
    (AMBER, "3", "Schedule MOU meeting — African Olympiad Academy",
     "Book a formal meeting. Prepare one-page platform overview + demo script. Olympiad content is seeded and ready."),
    (AMBER, "4", "Kenyan IMO team demo session",
     "Run one live demo. Goal: at least 5 IMO students active on platform by end of week."),
]

y = y0 + Inches(0.42)
for color, num, title, desc in must_ship:
    # Number badge
    nb = s3.shapes.add_shape(9, col_left, y, Inches(0.32), Inches(0.32))
    nb.fill.solid()
    nb.fill.fore_color.rgb = color
    nb.line.fill.background()
    tf_nb = nb.text_frame
    p_nb = tf_nb.paragraphs[0]
    p_nb.alignment = PP_ALIGN.CENTER
    r_nb = p_nb.add_run()
    r_nb.text = num
    r_nb.font.size = Pt(11)
    r_nb.font.bold = True
    r_nb.font.color.rgb = WHITE

    # Title
    txb_t = s3.shapes.add_textbox(col_left + Inches(0.4), y - Inches(0.02), Inches(5.6), Inches(0.28))
    tf_t = txb_t.text_frame
    p_t = tf_t.paragraphs[0]
    r_t = p_t.add_run()
    r_t.text = title
    r_t.font.size = Pt(11)
    r_t.font.bold = True
    r_t.font.color.rgb = DARK_TEXT

    # Desc
    txb_d = s3.shapes.add_textbox(col_left + Inches(0.4), y + Inches(0.24), Inches(5.6), Inches(0.38))
    tf_d = txb_d.text_frame
    tf_d.word_wrap = True
    p_d = tf_d.paragraphs[0]
    r_d = p_d.add_run()
    r_d.text = desc
    r_d.font.size = Pt(8.5)
    r_d.font.color.rgb = SLATE
    r_d.font.italic = True

    y += Inches(0.77)

# ── RIGHT: Should Ship + Success criteria ─────────────────────
section_header(s3, "Should Ship", col_right, y0, Inches(6.1), color=AMBER, font_size=13)

should_ship = [
    (AMBER, "5", "Personalisation prompts via notifications",
     "3-5 timed prompts asking users their goals, subjects, and level. Non-blocking notification badges, not pop-ups."),
    (SLATE, "6", "Guest trial conversion tracking",
     "Monitor guest-to-signup conversion rate. Set up day-7 nudge notification prompting account creation."),
]

y = y0 + Inches(0.42)
for color, num, title, desc in should_ship:
    nb = s3.shapes.add_shape(9, col_right, y, Inches(0.32), Inches(0.32))
    nb.fill.solid()
    nb.fill.fore_color.rgb = color
    nb.line.fill.background()
    tf_nb = nb.text_frame
    p_nb = tf_nb.paragraphs[0]
    p_nb.alignment = PP_ALIGN.CENTER
    r_nb = p_nb.add_run()
    r_nb.text = num
    r_nb.font.size = Pt(11)
    r_nb.font.bold = True
    r_nb.font.color.rgb = WHITE

    txb_t = s3.shapes.add_textbox(col_right + Inches(0.4), y - Inches(0.02), Inches(5.6), Inches(0.28))
    tf_t = txb_t.text_frame
    p_t = tf_t.paragraphs[0]
    r_t = p_t.add_run()
    r_t.text = title
    r_t.font.size = Pt(11)
    r_t.font.bold = True
    r_t.font.color.rgb = DARK_TEXT

    txb_d = s3.shapes.add_textbox(col_right + Inches(0.4), y + Inches(0.24), Inches(5.6), Inches(0.38))
    tf_d = txb_d.text_frame
    tf_d.word_wrap = True
    p_d = tf_d.paragraphs[0]
    r_d = p_d.add_run()
    r_d.text = desc
    r_d.font.size = Pt(8.5)
    r_d.font.color.rgb = SLATE
    r_d.font.italic = True

    y += Inches(0.77)

# ── Success criteria box ──────────────────────────────────────
y_box = Inches(3.55)
box_h = Inches(2.85)
success_box = s3.shapes.add_shape(5, col_right, y_box, Inches(6.0), box_h)
success_box.fill.solid()
success_box.fill.fore_color.rgb = EMERALD_LIGHT
success_box.line.color.rgb = EMERALD
success_box.line.width = Pt(1.2)

txb_hdr = s3.shapes.add_textbox(col_right + Inches(0.2), y_box + Inches(0.15),
                                  Inches(5.6), Inches(0.3))
tf_h = txb_hdr.text_frame
p_h = tf_h.paragraphs[0]
r_h = p_h.add_run()
r_h.text = "Week 2 — Success Looks Like"
r_h.font.size = Pt(11)
r_h.font.bold = True
r_h.font.color.rgb = EMERALD

criteria = [
    "At least 1 paying customer",
    "Plan selection modal live on sign-up",
    "MOU meeting booked with African Olympiad Academy",
    "5+ Kenyan IMO students active on the platform",
    "Guest trial conversion tracking active",
]
y_cr = y_box + Inches(0.52)
for c in criteria:
    # Check
    chk = s3.shapes.add_shape(9, col_right + Inches(0.18), y_cr + Inches(0.04),
                               Inches(0.18), Inches(0.18))
    chk.fill.solid()
    chk.fill.fore_color.rgb = EMERALD
    chk.line.fill.background()
    tf_chk = chk.text_frame
    p_chk = tf_chk.paragraphs[0]
    p_chk.alignment = PP_ALIGN.CENTER
    r_chk = p_chk.add_run()
    r_chk.text = "✓"
    r_chk.font.size = Pt(7)
    r_chk.font.bold = True
    r_chk.font.color.rgb = WHITE

    txb_c = s3.shapes.add_textbox(col_right + Inches(0.44), y_cr, Inches(5.4), Inches(0.28))
    tf_c = txb_c.text_frame
    p_c = tf_c.paragraphs[0]
    r_c = p_c.add_run()
    r_c.text = c
    r_c.font.size = Pt(10)
    r_c.font.color.rgb = EMERALD
    y_cr += Inches(0.44)

# Divider
div3 = s3.shapes.add_shape(1, Inches(6.7), Inches(1.0), Pt(1), Inches(6.0))
div3.fill.solid()
div3.fill.fore_color.rgb = RGBColor(0xC7, 0xD2, 0xFE)
div3.line.fill.background()

# Footer
bg_rect(s3, INDIGO_DARK, left=0, top=Inches(7.18), width=prs.slide_width, height=Inches(0.32))
add_textbox(s3, "EduReach  |  Week 2 Plan  |  May 2026",
            Inches(0.4), Inches(7.2), Inches(7), Inches(0.25),
            font_size=8, color=RGBColor(0xA5, 0xB4, 0xFC))
add_textbox(s3, "Slide 3 / 3",
            Inches(11.5), Inches(7.2), Inches(1.5), Inches(0.25),
            font_size=8, color=RGBColor(0xA5, 0xB4, 0xFC), align=PP_ALIGN.RIGHT)


# ══════════════════════════════════════════════════════════════════
prs.save("slides.pptx")
print("slides.pptx generated successfully.")
